
import streamlit as st
import pandas as pd
import numpy as np
import datetime
import time
import os
import gspread
from google.oauth2.service_account import Credentials
import streamlit.components.v1 as components
import pytz
import io
import requests
import zipfile
import traceback
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import black
from reportlab.lib.enums import TA_CENTER, TA_RIGHT, TA_LEFT

# Optional: Register a custom TTF font (e.g., Great Vibes for cursive effect)
# Place TTF file in your repo, e.g., db/fonts/GreatVibes-Regular.ttf
# Download from: https://fonts.google.com/specimen/Great+Vibes
# Uncomment and adjust path if you include the TTF file
# try:
#     pdfmetrics.registerFont(TTFont('GreatVibes', os.path.join(BASE_DIR, 'db/fonts/GreatVibes-Regular.ttf')))
# except Exception as e:
#     st.warning(f"Failed to load custom font: {str(e)}. Using Helvetica instead.")

def get_template_path(template_type):
    template_path = os.path.join(DB_FOLDER, f"{template_type}_template.docx")
    if os.path.exists(template_path):
        return template_path
    
    github_url = f"https://raw.githubusercontent.com/your_username/your_repo_name/main/db/{template_type}_template.docx"
    try:
        response = requests.get(github_url, timeout=10)
        if response.status_code == 200:
            temp_path = f"/tmp/{template_type}_template.docx"
            with open(temp_path, "wb") as file:
                file.write(response.content)
            return temp_path
        else:
            raise Exception(f"HTTP {response.status_code}")
    except Exception as e:
        st.error(f"Failed to download {template_type} template from GitHub: {str(e)}. Please add 'db/{template_type}_template.docx' to your repo.")
        return None

def generate_certificate(emp_id, emp_name, test_date, status, template_type):
    template_path = get_template_path(template_type)
    if not template_path:
        st.error(f"No {template_type} template available. Cannot generate certificate.")
        return None, None

    try:
        # Parse DOCX to extract content and styling
        doc = Document(template_path)
        date_str = test_date.split()[0] if " " in test_date else test_date
        try:
            test_date_obj = datetime.datetime.strptime(date_str, "%d-%m-%Y")
        except ValueError:
            st.error(f"Invalid date format in test_date: {test_date}. Expected format: DD-MM-YYYY")
            return None, None
        validity_date_obj = test_date_obj + datetime.timedelta(days=5*365)
        cert_number = f"{emp_id}/PTIS/{template_type}/{date_str.replace('-', '')}"
        status_text = 'Pass' if status == "Pass" else 'Fail'

        # Create PDF buffer
        buffer = io.BytesIO()
        c = canvas.Canvas(buffer, pagesize=A4)
        width, height = A4

        # Process DOCX paragraphs and map to PDF
        y_position = height - inch  # Start near top of A4
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            # Get alignment and font settings from DOCX
            alignment = para.alignment if para.alignment else WD_ALIGN_PARAGRAPH.LEFT
            font_name = 'Helvetica'  # Default fallback
            font_size = 12
            for run in para.runs:
                if run.font.name and 'Corsiva' in run.font.name:
                    font_name = 'GreatVibes' if 'GreatVibes' in pdfmetrics.getRegisteredFontNames() else 'Helvetica-Bold'
                if run.font.size:
                    font_size = run.font.size.pt
                break  # Use first run's style

            # Replace placeholders
            if 'Usman Waheed' in text:
                text = text.replace('Usman Waheed', emp_name)
                font_name = 'GreatVibes' if 'GreatVibes' in pdfmetrics.getRegisteredFontNames() else 'Helvetica-Bold'
                font_size = 26
                align = TA_CENTER
            elif '25-September-2025' in text or 'Date of Certification' in text:
                text = text.replace('25-September-2025', test_date_obj.strftime("%d-%B-%Y"))
                align = TA_RIGHT if template_type != "MT" else TA_CENTER
                if template_type == "MT":
                    text = text + "            "  # Simulate padding
            elif '25/PTIS/DPT/00410' in text:
                text = text.replace('25/PTIS/DPT/00410', cert_number)
                align = TA_LEFT if template_type not in ["MT", "VT"] else TA_CENTER if template_type == "MT" else TA_LEFT
                if template_type == "MT":
                    text = "            " + text  # Simulate padding
                elif template_type == "VT":
                    text = "  " + text
            elif 'Validity: 24-September-2030' in text:
                text = text.replace('Validity: 24-September-2030', f'Validity: {validity_date_obj.strftime("%d-%B-%Y")}')
                align = TA_RIGHT if template_type != "MT" else TA_CENTER
                if template_type == "MT":
                    text = text + "            "  # Simulate padding
            elif 'Status' in text:
                text = text.replace('Status: Fail', status_text).replace('Status: Pass', status_text)
                align = TA_LEFT
            else:
                align = {WD_ALIGN_PARAGRAPH.CENTER: TA_CENTER, WD_ALIGN_PARAGRAPH.RIGHT: TA_RIGHT, WD_ALIGN_PARAGRAPH.LEFT: TA_LEFT}.get(alignment, TA_LEFT)

            # Draw text on PDF
            c.setFont(font_name, font_size)
            text_width = c.stringWidth(text, font_name, font_size)
            if align == TA_CENTER:
                x = width / 2
            elif align == TA_RIGHT:
                x = width - inch - text_width
            else:
                x = inch
            c.drawString(x, y_position, text)
            y_position -= font_size + 10  # Adjust spacing

        c.showPage()
        c.save()
        buffer.seek(0)

        safe_name = "".join(c for c in emp_name if c.isalnum() or c in (' ', '-', '_')).rstrip()
        certificate_filename = f"{template_type}_Certificate_{emp_id}_{safe_name}_{date_str}.pdf"
        st.success(f"Generated PDF certificate: {certificate_filename}")
        return buffer.getvalue(), certificate_filename

    except Exception as e:
        st.error(f"Error generating {template_type} certificate: {str(e)}")
        return None, None

# Certificate Generation Section (replace in your main code under "Generate Certificates")
st.markdown("---")
st.subheader("📜 Generate Certificates")

# Simple certificate filter - only employee name
results_df = load_all_results()  # Ensure results_df is loaded
passed_results = results_df[results_df["Status"] == "Pass"]
cert_employee_names = ["All"] + sorted(passed_results["Name"].unique().tolist())

selected_cert_name = st.selectbox(
    "Filter Certificates by Employee Name",
    cert_employee_names,
    index=0,
    key=f"cert_name_filter_{st.session_state.filter_reset_counter}"
)

if st.button("Generate Certificates for Qualifying Employees"):
    required_standards = {"DS-1", "Cummulative", "API SPEC 5CT & 5A5", "API RP 7G-2"}
    passed_results = results_df[results_df["Status"] == "Pass"]
    
    # Initialize qualifying_rows
    qualifying_rows = []
    if not passed_results.empty:
        grouped = passed_results.groupby('Name')
        for name, group in grouped:
            passed_standards = set(group['Test Type'].str.strip())
            if required_standards.issubset(passed_standards):
                cumm_row = group[group['Test Type'].str.strip() == 'Cummulative']
                if not cumm_row.empty:
                    qualifying_rows.append(cumm_row.iloc[0])
    
    # Create qualifying_df
    qualifying_df = pd.DataFrame(qualifying_rows)
    
    # Apply name filter
    if selected_cert_name != "All":
        qualifying_df = qualifying_df[qualifying_df["Name"] == selected_cert_name]
    
    # Check if qualifying_df is empty
    if qualifying_df.empty:
        st.warning("No qualifying candidates found. Ensure employees have passed all required standards (DS-1, Cummulative, API SPEC 5CT & 5A5, API RP 7G-2).")
    else:
        st.info(f"Found {len(qualifying_df)} qualifying candidate(s) for certificate generation.")
        certificate_files = []
        for _, row in qualifying_df.iterrows():
            emp_id = str(row['ID'])
            emp_name = str(row['Name'])
            test_date = str(row['Date / Time'])
            status = str(row['Status'])
            for template_type in ['PT', 'UT', 'MT', 'VT']:
                certificate_data, certificate_filename = generate_certificate(emp_id, emp_name, test_date, status, template_type)
                if certificate_data:
                    certificate_files.append((certificate_data, certificate_filename))
        
        if certificate_files:
            zip_buffer = io.BytesIO()
            with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zipf:
                for cert_data, cert_filename in certificate_files:
                    zipf.writestr(cert_filename, cert_data)  # Write PDF data directly
            zip_buffer.seek(0)
            filename_suffix = selected_cert_name if selected_cert_name != "All" else "all_qualifying"
            st.download_button(
                label=f"Download Certificates (ZIP) for {filename_suffix}",
                data=zip_buffer,
                file_name=f"certificates_{filename_suffix}_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.zip",
                mime="application/zip"
            )
        else:
            st.error("Failed to generate any certificates. Check templates and permissions.")
        
            st.markdown("---")
            if st.button("Logout"):
                st.session_state.admin_logged_in = False
                st.session_state.pop("quiz", None)
                st.rerun()
            else:
                st.info("No results available yet in the Result 2 sheet.")
                if st.button("Logout"):
                    st.session_state.admin_logged_in = False
                    st.session_state.pop("quiz", None)
                    st.rerun()

# Employee login and quiz
if not st.session_state.admin_logged_in and "quiz" not in st.session_state:
    st.subheader("Employee Login")
    col1, col2 = st.columns(2)
    with col1:
        emp_id = st.text_input(
            "Employee ID", 
            value="", 
            key=f"id_{st.session_state.reset_counter}",
            help="Enter your employee identification number"
        )
    def fetch_name(employees_df, emp_id_input):
        if emp_id_input and not employees_df.empty:
            try:
                fetched = employees_df[employees_df["ID"].astype(str).str.strip() == str(emp_id_input).strip()]
                if not fetched.empty:
                    return str(fetched.iloc[0]["Name"])
            except Exception:
                pass
        return ""
    fetched_name = fetch_name(employees, emp_id) if "name" not in st.session_state else st.session_state["name"]
    with col2:
        name = st.text_input(
            "Name", 
            value=fetched_name, 
            key=f"name_{st.session_state.reset_counter}",
            help="This will auto-fill if your Employee ID is found"
        )
    options = standards["Standard"].dropna().unique().tolist()
    options = sorted(options)
    if "Cummulative" not in options:
        options = ["Cummulative"] + options
    selected_standard = st.selectbox("Select Standard", options, index=0 if options else None, key=f"std_{st.session_state.reset_counter}")
    total, criteria, h, m, s = get_info_for_standard(standards, selected_standard)
    c1, c2, c3 = st.columns(3)
    with c1: st.metric("Total Questions", total)
    with c2: st.metric("Passing Criteria (%)", criteria)
    with c3: st.metric("Timer (HH:MM:SS)", f"{h:02d}:{m:02d}:{s:02d}")
    st.markdown("---")
    with st.form("start_form"):
        st.markdown("### Ready to start your test?")
        col1, col2, col3 = st.columns([1, 2, 1])
        with col2:
            submitted = st.form_submit_button("🚀 Start Test", use_container_width=True)
    if submitted:
        if not emp_id or not name or not selected_standard:
            st.error("Please enter ID, Name and select a Standard.")
        else:
            ok, msg = start_quiz_session(emp_id, name, selected_standard, questions, total)
            if not ok:
                st.error(msg)
            else:
                st.rerun()

# Quiz interface
elif "quiz" in st.session_state:
    qstate = st.session_state.quiz
    total, criteria, h, m, s = get_info_for_standard(standards, qstate["standard"])
    total_secs = format_timer(h, m, s)

    elapsed = int(time.time() - qstate["start_ts"])
    remaining = max(0, total_secs - elapsed)

    if total_secs > 0 and len(qstate["queue"]) > 0 and "submitted" not in st.session_state:
        if remaining <= 0:
            st.error("Time is up! Auto-submitting your test...")
            for qid in qstate["queue"]:
                if qid not in qstate.get("attempted", set()):
                    qstate["wrong"] += 1
            qstate["queue"] = []
            st.session_state.quiz = qstate
            
            right, wrong, total_q = qstate["right"], qstate["wrong"], qstate["total"]
            raw_score = right - (wrong * 0.25)
            final_score = max(0, raw_score)
            pct = (final_score/total_q)*100 if total_q else 0.0
            
            status = "Pass" if pct >= float(criteria) else "Fail"
            ok, msg = append_result(
                qstate["emp_id"], qstate["emp_name"], total_q, right, wrong, criteria, status, qstate["standard"]
            )
            st.session_state["submitted"] = True
            st.session_state["submit_result"] = (ok, msg, right, wrong, total_q, pct, criteria, status, final_score)
            st.query_params.clear()
            st.rerun()

        rem_h = remaining // 3600
        rem_m = (remaining % 3600) // 60
        rem_s = remaining % 60

        if remaining <= 300:
            bg_color = "#DC2626"
            text_color = "white"
            icon = "🚨"
            pulse_class = "timer-pulse"
        elif remaining <= 900:
            bg_color = "#DC2626"
            text_color = "white"
            icon = "⚠️"
            pulse_class = ""
        elif remaining <= 1800:
            bg_color = "#D97706"
            text_color = "white"
            icon = "⏰"
            pulse_class = ""
        else:
            bg_color = "#1E3A8A"
            text_color = "white"
            icon = "⏰"
            pulse_class = ""

        progress_percent = (remaining / total_secs) * 100 if total_secs > 0 else 0

        timer_html = f"""
        <style>
        @keyframes pulse {{
            0% {{ transform: scale(1); opacity: 1; }}
            50% {{ transform: scale(1.05); opacity: 0.8; }}
            100% {{ transform: scale(1); opacity: 1; }}
        }}
        .timer-pulse {{
            animation: pulse 1s infinite;
        }}
        .timer-container {{
            padding: 20px;
            border-radius: 15px;
            text-align: center;
            font-size: 22px;
            font-weight: bold;
            margin-bottom: 20px;
            box-shadow: 0 8px 16px rgba(0, 0, 0, 0.2);
            border: 3px solid rgba(255, 255, 255, 0.1);
        }}
        </style>
        <div id="timer_container" class="timer-container {pulse_class}" style="background: linear-gradient(135deg, {bg_color}, {bg_color}CC); color: {text_color};">
            <div style="display: flex; align-items: center; justify-content: center; gap: 15px;">
                <span id="timer_icon" style="font-size: 28px;">{icon}</span>
                <span>Time Remaining :</span>
                <span id="timer_display" style="font-family: 'Courier New', monospace; font-size: 28px; background: rgba(0,0,0,0.2); padding: 5px 15px; border-radius: 8px;">
                    {rem_h:02d}:{rem_m:02d}:{rem_s:02d}
                </span>
            </div>
            <div style="width: 100%; height: 6px; background-color: rgba(255,255,255,0.3); border-radius: 3px; overflow: hidden; margin-top: 15px;">
                <div id="progress_bar" style="height: 100%; background: linear-gradient(90deg, #10B981, #34D399); width: {progress_percent:.1f}%; border-radius: 3px; transition: width 0.5s ease-in-out;"></div>
            </div>
        </div>
        <script>
        (function() {{
            var remaining = {remaining};
            var total_secs = {total_secs};
            var interval = null;

            function updateTimer() {{
                if (remaining <= 0) {{
                    document.getElementById('timer_display').innerText = '00:00:00';
                    document.getElementById('progress_bar').style.width = '0%';
                    clearInterval(interval);
                    var form = document.createElement('form');
                    form.method = 'POST';
                    form.action = window.location.href;
                    var input = document.createElement('input');
                    input.type = 'hidden';
                    input.name = 'timeout';
                    input.value = 'true';
                    form.appendChild(input);
                    document.body.appendChild(form);
                    form.submit();
                    return;
                }}
                var h = Math.floor(remaining / 3600);
                var m = Math.floor((remaining % 3600) / 60);
                var s = remaining % 60;
                document.getElementById('timer_display').innerText = `${{h.toString().padStart(2, '0')}}:${{m.toString().padStart(2, '0')}}:${{s.toString().padStart(2, '0')}}`;
                var progress = (remaining / total_secs) * 100;
                document.getElementById('progress_bar').style.width = progress + '%';
                var container = document.getElementById('timer_container');
                var iconElem = document.getElementById('timer_icon');
                var bg_color, text_color, icon, pulse_class = '';
                if (remaining <= 300) {{
                    bg_color = '#DC2626';
                    text_color = 'white';
                    icon = '🚨';
                    pulse_class = 'timer-pulse';
                }} else if (remaining <= 900) {{
                    bg_color = '#DC2626';
                    text_color = 'white';
                    icon = '⚠️';
                }} else if (remaining <= 1800) {{
                    bg_color = '#D97706';
                    text_color = 'white';
                    icon = '⏰';
                }} else {{
                    bg_color = '#1E3A8A';
                    text_color = 'white';
                    icon = '⏰';
                }}
                container.style.background = `linear-gradient(135deg, ${bg_color}, ${bg_color}CC)`;
                container.style.color = text_color;
                iconElem.innerText = icon;
                if (pulse_class) {{
                    container.classList.add(pulse_class);
                }} else {{
                    container.classList.remove('timer-pulse');
                }}
                remaining--;
            }}

            if (interval) {{
                clearInterval(interval);
            }}
            updateTimer();
            interval = setInterval(updateTimer, 1000);
        }})();
        </script>
        """
        components.html(timer_html, height=150)

        if st.query_params.get("timeout", ["false"])[0] == "true":
            if len(qstate["queue"]) > 0:
                st.error("Time is up! Auto-submitting your test...")
                for qid in qstate["queue"]:
                    if qid not in qstate.get("attempted", set()):
                        qstate["wrong"] += 1
                qstate["queue"] = []
                st.session_state.quiz = qstate
                
                right, wrong, total_q = qstate["right"], qstate["wrong"], qstate["total"]
                raw_score = right - (wrong * 0.25)
                final_score = max(0, raw_score)
                pct = (final_score/total_q)*100 if total_q else 0.0
                
                status = "Pass" if pct >= float(criteria) else "Fail"
                ok, msg = append_result(
                    qstate["emp_id"], qstate["emp_name"], total_q, right, wrong, criteria, status, qstate["standard"]
                )
                st.session_state["submitted"] = True
                st.session_state["submit_result"] = (ok, msg, right, wrong, total_q, pct, criteria, status, final_score)
                st.query_params.clear()
                st.rerun()

        if remaining <= 300:
            st.markdown('<div style="margin-bottom: 8px;"></div>', unsafe_allow_html=True)
            st.warning("🚨 URGENT: Less than 5 minutes remaining!")
        elif remaining <= 900:
            st.markdown('<div style="margin-bottom: 8px;"></div>', unsafe_allow_html=True)
            st.warning("⚠️ WARNING: Less than 15 minutes remaining!")
        elif remaining <= 1800:
            st.markdown('<div style="margin-bottom: 8px;"></div>', unsafe_allow_html=True)
            st.info("⏰ NOTICE: Less than 30 minutes remaining!")

    elif total_secs > 0 and "submitted" in st.session_state:
        rem_h = remaining // 3600
        rem_m = (remaining % 3600) // 60
        rem_s = remaining % 60

        if remaining <= 300:
            bg_color = "#DC2626"
            text_color = "white"
            icon = "🚨"
            pulse_class = "timer-pulse"
        elif remaining <= 900:
            bg_color = "#DC2626"
            text_color = "white"
            icon = "⚠️"
            pulse_class = ""
        elif remaining <= 1800:
            bg_color = "#D97706"
            text_color = "white"
            icon = "⏰"
            pulse_class = ""
        else:
            bg_color = "#1E3A8A"
            text_color = "white"
            icon = "⏰"
            pulse_class = ""

        progress_percent = (remaining / total_secs) * 100 if total_secs > 0 else 0

        stopped_timer_html = f"""
        <style>
        @keyframes pulse {{
            0% {{ transform: scale(1); opacity: 1; }}
            50% {{ transform: scale(1.05); opacity: 0.8; }}
            100% {{ transform: scale(1); opacity: 1; }}
        }}
        .timer-pulse {{
            animation: pulse 1s infinite;
        }}
        .timer-container {{
            padding: 20px;
            border-radius: 15px;
            text-align: center;
            font-size: 22px;
            font-weight: bold;
            margin-bottom: 20px;
            box-shadow: 0 8px 16px rgba(0, 0, 0, 0.2);
            border: 3px solid rgba(255, 255, 255, 0.1);
        }}
        </style>
        <div id="timer_container" class="timer-container {pulse_class}" style="background: linear-gradient(135deg, {bg_color}, {bg_color}CC); color: {text_color};">
            <div style="display: flex; align-items: center; justify-content: center; gap: 15px;">
                <span id="timer_icon" style="font-size: 28px;">{icon}</span>
                <span>Test Submitted - Time Remaining :</span>
                <span id="timer_display" style="font-family: 'Courier New', monospace; font-size: 28px; background: rgba(0,0,0,0.2); padding: 5px 15px; border-radius: 8px;">
                    {rem_h:02d}:{rem_m:02d}:{rem_s:02d}
                </span>
            </div>
            <div style="width: 100%; height: 6px; background-color: rgba(255,255,255,0.3); border-radius: 3px; overflow: hidden; margin-top: 15px;">
                <div id="progress_bar" style="height: 100%; background: linear-gradient(90deg, #10B981, #34D399); width: {progress_percent:.1f}%; border-radius: 3px;"></div>
            </div>
        </div>
        """
        components.html(stopped_timer_html, height=150)

    if "attempted" not in st.session_state.quiz:
        st.session_state.quiz["attempted"] = set()
    if "skipped_questions" not in st.session_state.quiz:
        st.session_state.quiz["skipped_questions"] = set()

    answered_count = qstate["total"] - len(qstate["queue"])

    st.markdown(
        f"""
        <div style="padding: 12px 15px; border-radius: 8px; background: linear-gradient(135deg, #1E3A8A, #3B82F6); color: white; text-align: center; font-size: 17px; margin-bottom: 20px; white-space: nowrap; overflow: hidden; box-shadow: 0 2px 4px rgba(0,0,0,0.1);">
            <b>ID :</b> {qstate['emp_id']} &nbsp;•&nbsp; <b>Name :</b> {qstate['emp_name']} &nbsp;•&nbsp; <b>Standard :</b> {qstate['standard']} &nbsp;•&nbsp; <b>Progress :</b> {answered_count}/{qstate['total']}
        </div>
        """,
        unsafe_allow_html=True
    )

    st.info("📌 **Scoring System**: +1 mark for correct answer, -0.25 marks for wrong answer, 0 marks for unattempted questions")

    if len(qstate["queue"]) > 0:
        current_qid = qstate["queue"][0]
        row = qstate["rows"].iloc[current_qid]
        qno, question, A, B, C, D, correct = row["Qno"], row["Question"], row["A"], row["B"], row["C"], row["D"], row["Answer"]

        is_previously_skipped = current_qid in qstate["skipped_questions"]
        
        if is_previously_skipped:
            st.markdown("🔄 **This question was skipped earlier**")
            st.subheader(f"Q{current_qid+1}. {question}")
        else:
            st.subheader(f"Q{current_qid+1}. {question}")
            
        choice = st.radio("Choose your answer:", [A, B, C, D], index=None, key=f"q_{current_qid}")

        col1, col2 = st.columns([1,1])

        with col1:
            if st.button("Next", use_container_width=True):
                if choice is None:
                    st.warning("⚠️ Please select an option before moving on.")
                else:
                    qstate["attempted"].add(current_qid)
                    mapping = {"A": A, "B": B, "C": C, "D": D}
                    correct_text = mapping.get(str(correct).strip(), str(correct).strip())
                    is_correct = str(choice).strip() == str(correct_text).strip()
                    qstate["answers"][current_qid] = {
                        "choice": choice,
                        "correct": correct_text,
                        "is_correct": is_correct
                    }
                    if is_correct:
                        qstate["right"] += 1
                    else:
                        qstate["wrong"] += 1
                    qstate["queue"].pop(0)
                    st.session_state.quiz = qstate
                    st.rerun()

        with col2:
            if len(qstate["queue"]) > 1 and not is_previously_skipped:
                if st.button("Skip", use_container_width=True):
                    qstate["skipped_questions"].add(current_qid)
                    qstate["queue"].append(qstate["queue"].pop(0))
                    st.session_state.quiz = qstate
                    st.rerun()

    if len(qstate["queue"]) == 0 and "submitted" not in st.session_state:
        right, wrong, total_q = qstate["right"], qstate["wrong"], qstate["total"]
        raw_score = right - (wrong * 0.25)
        final_score = max(0, raw_score)
        pct = (final_score/total_q)*100 if total_q else 0.0
        status = "Pass" if pct >= float(criteria) else "Fail"

        st.success("All questions attempted. You can now submit your test.")

        submit_clicked = st.button("Submit", use_container_width=True)
        if submit_clicked:
            ok, msg = append_result(
                qstate["emp_id"], qstate["emp_name"], total_q, right, wrong, criteria, status, qstate["standard"]
            )
            st.session_state["submitted"] = True
            st.session_state["submit_result"] = (ok, msg, right, wrong, total_q, pct, criteria, status, final_score)
            st.rerun()

    if "submitted" in st.session_state:
        if "submit_result" in st.session_state:
            result_data = st.session_state["submit_result"]
            ok, msg, right, wrong, total_q, pct, criteria, status, final_score = result_data
            
            if not ok:
                st.error(f"Failed to save results to Google Sheets: {msg}")

            color = "#10B981" if status == "Pass" else "#DC2626"
            st.markdown(
                f"""
                <div style="padding:20px; border-radius:12px; background: linear-gradient(135deg, #1E3A8A, #2563EB); color:white; text-align:center; margin-top:20px;">
                    <h3 style="color:{color}; font-weight:700;">Final Result : <span style="font-weight:700;">{status}</span></h3>
                    <p style="font-size:18px;">
                        <b>Correct Answers :</b> {right}<br>
                        <b>Wrong Answers :</b> {wrong}<br>
                        <b>Final Score :</b> {final_score:.2f}/{total_q}<br>
                        <b>Percentage :</b> {pct:.2f}%<br>
                        <b>Passing Criteria :</b> {criteria:.0f}%
                    </p>
                    <small style="opacity: 0.8;">Negative marking: -0.25 marks per wrong answer</small>
                </div>
                """,
                unsafe_allow_html=True
            )
